import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import random
from openpyxl import Workbook


def parse_float(value: str) -> float:
    """Accepte le point ou la virgule comme séparateur décimal."""
    return float(value.strip().replace(",", "."))


def generer():
    try:
        quantite = int(entree_quantite.get().strip())
        moyenne = parse_float(entree_moyenne.get())
        ecart_type = parse_float(entree_ecart_type.get())

        if quantite <= 0:
            raise ValueError("La quantité doit être strictement positive.")
        if ecart_type < 0:
            raise ValueError("L'écart-type ne peut pas être négatif.")

    except ValueError as exc:
        messagebox.showerror(
            "Paramètres invalides",
            f"Vérifie les valeurs saisies.\n\n{exc}"
        )
        return

    chemin = filedialog.asksaveasfilename(
        title="Enregistrer le fichier Excel",
        defaultextension=".xlsx",
        filetypes=[("Fichier Excel", "*.xlsx")],
        initialfile="nombres_loi_normale.xlsx",
    )

    if not chemin:
        return

    # Génération selon N(moyenne, écart-type²)
    valeurs = [random.gauss(moyenne, ecart_type) for _ in range(quantite)]

    wb = Workbook()
    ws = wb.active
    ws.title = "Loi normale"

    # Paramètres rappelés en haut du fichier
    ws["A1"] = "Paramètre"
    ws["B1"] = "Valeur"
    ws["A2"] = "Quantité"
    ws["B2"] = quantite
    ws["A3"] = "Moyenne demandée"
    ws["B3"] = moyenne
    ws["A4"] = "Écart-type demandé"
    ws["B4"] = ecart_type

    # Valeurs générées
    ws["A6"] = "Valeur"
    for ligne, valeur in enumerate(valeurs, start=7):
        ws.cell(row=ligne, column=1, value=valeur)

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 20

    wb.save(chemin)

    moyenne_observee = sum(valeurs) / len(valeurs)
    if len(valeurs) > 1:
        variance = sum((x - moyenne_observee) ** 2 for x in valeurs) / (len(valeurs) - 1)
        ecart_type_observe = variance ** 0.5
    else:
        ecart_type_observe = 0.0

    messagebox.showinfo(
        "Fichier généré",
        "Le fichier Excel a été créé avec succès.\n\n"
        f"Quantité : {quantite}\n"
        f"Moyenne observée : {moyenne_observee:.4f}\n"
        f"Écart-type observé : {ecart_type_observe:.4f}"
    )


root = tk.Tk()
root.title("Générateur de nombres — Loi normale")
root.resizable(False, False)

cadre = ttk.Frame(root, padding=20)
cadre.grid(row=0, column=0)

ttk.Label(
    cadre,
    text="Générateur selon une loi normale",
    font=("Segoe UI", 14, "bold")
).grid(row=0, column=0, columnspan=2, pady=(0, 18))

ttk.Label(cadre, text="Quantité de nombres :").grid(
    row=1, column=0, sticky="w", padx=(0, 15), pady=7
)
entree_quantite = ttk.Entry(cadre, width=22)
entree_quantite.insert(0, "1000")
entree_quantite.grid(row=1, column=1, pady=7)

ttk.Label(cadre, text="Moyenne :").grid(
    row=2, column=0, sticky="w", padx=(0, 15), pady=7
)
entree_moyenne = ttk.Entry(cadre, width=22)
entree_moyenne.insert(0, "50")
entree_moyenne.grid(row=2, column=1, pady=7)

ttk.Label(cadre, text="Écart-type :").grid(
    row=3, column=0, sticky="w", padx=(0, 15), pady=7
)
entree_ecart_type = ttk.Entry(cadre, width=22)
entree_ecart_type.insert(0, "10")
entree_ecart_type.grid(row=3, column=1, pady=7)

ttk.Button(
    cadre,
    text="Générer le fichier Excel",
    command=generer
).grid(row=4, column=0, columnspan=2, pady=(18, 0), ipadx=20, ipady=5)

root.mainloop()
