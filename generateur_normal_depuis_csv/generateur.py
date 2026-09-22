from __future__ import annotations

import csv
import math
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


INPUT_FILE = "input.csv"
OUTPUT_FILE = "output.xlsx"

NB_COLONNES_GENEREES = 10
ECART_TYPE = 10.0
MIN_VALUE = 0.0
MAX_VALUE = 100.0
DECIMALES = 2


def parse_number(value: str) -> float:
    value = value.strip().replace("\u00a0", "").replace(" ", "")
    value = value.replace(",", ".")
    number = float(value)

    if not math.isfinite(number):
        raise ValueError("la valeur n'est pas un nombre fini")

    return number


def detect_delimiter(text: str) -> str:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        return dialect.delimiter
    except csv.Error:
        return ";"


def read_input(path: Path) -> tuple[str, list[float]]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        text = path.read_text(encoding="cp1252")

    if not text.strip():
        raise ValueError("input.csv est vide.")

    delimiter = detect_delimiter(text)
    rows = list(csv.reader(text.splitlines(), delimiter=delimiter))
    rows = [row for row in rows if row and any(cell.strip() for cell in row)]

    if not rows:
        raise ValueError("input.csv ne contient aucune donnée.")

    first_cell = rows[0][0].strip()

    try:
        parse_number(first_cell)
        has_header = False
    except ValueError:
        has_header = True

    header = first_cell if has_header and first_cell else "valeur"
    data_rows = rows[1:] if has_header else rows

    values: list[float] = []

    for line_number, row in enumerate(data_rows, start=2 if has_header else 1):
        if not row or not row[0].strip():
            continue

        try:
            value = parse_number(row[0])
        except ValueError as exc:
            raise ValueError(
                f"Ligne {line_number} : '{row[0]}' n'est pas une valeur numérique valide."
            ) from exc

        if not MIN_VALUE <= value <= MAX_VALUE:
            raise ValueError(
                f"Ligne {line_number} : la valeur {value} est hors de l'intervalle "
                f"[{MIN_VALUE:g}, {MAX_VALUE:g}]."
            )

        values.append(value)

    if not values:
        raise ValueError("Aucune valeur numérique n'a été trouvée dans la première colonne.")

    return header, values


def truncated_normal(mean: float) -> float:
    while True:
        value = random.gauss(mean, ECART_TYPE)
        if MIN_VALUE <= value <= MAX_VALUE:
            return round(value, DECIMALES)


def create_output(source_header: str, source_values: list[float], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Données"

    headers = [source_header] + [
        f"generation_{i}" for i in range(1, NB_COLONNES_GENEREES + 1)
    ]
    ws.append(headers)

    for source_value in source_values:
        generated = [
            truncated_normal(source_value)
            for _ in range(NB_COLONNES_GENEREES)
        ]
        ws.append([source_value] + generated)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col_idx in range(1, NB_COLONNES_GENEREES + 2):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 16
        for cell in ws[col_letter][1:]:
            cell.number_format = "0.00"

    wb.save(output_path)


def main() -> int:
    base_dir = Path(__file__).resolve().parent
    input_path = base_dir / INPUT_FILE
    output_path = base_dir / OUTPUT_FILE

    if not input_path.exists():
        print(f"ERREUR : fichier introuvable : {input_path.name}")
        print("Place un fichier nommé 'input.csv' dans le même dossier que le programme.")
        return 1

    try:
        source_header, source_values = read_input(input_path)
        create_output(source_header, source_values, output_path)
    except Exception as exc:
        print(f"ERREUR : {exc}")
        return 1

    print()
    print("Génération terminée.")
    print(f"Nombre de lignes traitées : {len(source_values)}")
    print(f"Colonnes générées par ligne : {NB_COLONNES_GENEREES}")
    print(f"Écart-type : {ECART_TYPE:g}")
    print(f"Bornes : [{MIN_VALUE:g}, {MAX_VALUE:g}]")
    print(f"Fichier créé : {output_path}")
    print()
    return 0


if __name__ == "__main__":
    sys.exit(main())
