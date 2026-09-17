import csv
import secrets
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None

ALPHABET = "23456789abcdefghjklmnpqrstuvwxyz"
ID_LENGTH = 4
MAX_IDS = len(ALPHABET) ** ID_LENGTH

BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "id_input.csv"
OUTPUT_CSV = BASE_DIR / "id_output.csv"
OUTPUT_XLSX = BASE_DIR / "id_output.xlsx"


def is_valid_id(value: str) -> bool:
    return len(value) == ID_LENGTH and all(char in ALPHABET for char in value)


def read_existing_ids():
    if not INPUT_FILE.exists():
        return []

    with INPUT_FILE.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)

        if not reader.fieldnames or "id" not in reader.fieldnames:
            raise ValueError('Le fichier "id_input.csv" doit contenir une colonne nommée exactement "id".')

        ids = []
        seen = set()

        for line_number, row in enumerate(reader, start=2):
            raw_value = row.get("id", "")
            value = (raw_value or "").strip().lower()

            if not value:
                continue

            if not is_valid_id(value):
                raise ValueError(
                    f'ID invalide à la ligne {line_number} : "{value}". '
                    f"Chaque ID doit contenir exactement {ID_LENGTH} caractères parmi : {ALPHABET}"
                )

            if value in seen:
                raise ValueError(f'Doublon détecté à la ligne {line_number} : "{value}".')

            seen.add(value)
            ids.append(value)

        return ids


def generate_unique_id():
    return "".join(secrets.choice(ALPHABET) for _ in range(ID_LENGTH))


def write_csv(ids):
    with OUTPUT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["id"])
        for value in ids:
            writer.writerow([value])


def write_xlsx(ids):
    if Workbook is None:
        raise RuntimeError(
            "La bibliothèque openpyxl n'est pas installée. "
            "Lancez installer.bat une fois avant d'utiliser le programme."
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "IDs"

    ws["A1"] = "id"
    ws["A1"].number_format = "@"

    for row_index, value in enumerate(ids, start=2):
        cell = ws.cell(row=row_index, column=1, value=value)
        cell.number_format = "@"
        cell.data_type = "s"

    ws.column_dimensions["A"].width = 12
    wb.save(OUTPUT_XLSX)


def ask_total(existing_count):
    while True:
        raw = input(
            f"\nNombre total d'identifiants souhaité "
            f"(existants inclus, min. {existing_count}, max. {MAX_IDS}) : "
        ).strip()

        try:
            total = int(raw)
        except ValueError:
            print("Veuillez saisir un nombre entier.")
            continue

        if total < existing_count:
            print(f"Le total ne peut pas être inférieur au nombre d'IDs existants ({existing_count}).")
            continue

        if total > MAX_IDS:
            print(f"Le total ne peut pas dépasser {MAX_IDS:,} IDs possibles.".replace(",", " "))
            continue

        if total < 0:
            print("Le total ne peut pas être négatif.")
            continue

        return total


def main():
    print("=" * 62)
    print(" GENERATEUR D'IDENTIFIANTS UNIQUES")
    print("=" * 62)
    print(f"Format : {ID_LENGTH} caractères")
    print(f"Caractères autorisés : {ALPHABET}")
    print(f"Nombre maximal d'IDs possibles : {MAX_IDS:,}".replace(",", " "))

    try:
        existing_ids = read_existing_ids()
    except Exception as exc:
        print(f"\nERREUR : {exc}")
        return 1

    print(f"\nIDs existants trouvés : {len(existing_ids)}")

    total_needed = ask_total(len(existing_ids))
    new_needed = total_needed - len(existing_ids)

    if new_needed == 0:
        final_ids = list(existing_ids)
    else:
        used = set(existing_ids)
        new_ids = []

        if total_needed <= MAX_IDS // 2:
            while len(new_ids) < new_needed:
                candidate = generate_unique_id()
                if candidate not in used:
                    used.add(candidate)
                    new_ids.append(candidate)
        else:
            import itertools
            for chars in itertools.product(ALPHABET, repeat=ID_LENGTH):
                candidate = "".join(chars)
                if candidate not in used:
                    used.add(candidate)
                    new_ids.append(candidate)
                    if len(new_ids) >= new_needed:
                        break

        final_ids = existing_ids + new_ids

    try:
        write_csv(final_ids)
        write_xlsx(final_ids)
    except Exception as exc:
        print(f"\nERREUR lors de l'écriture des fichiers : {exc}")
        return 1

    print("\nTerminé.")
    print(f"IDs existants conservés : {len(existing_ids)}")
    print(f"Nouveaux IDs générés    : {new_needed}")
    print(f"Total                    : {len(final_ids)}")
    print("\nFichiers créés :")
    print(f"- {OUTPUT_CSV.name}  (CSV brut)")
    print(f"- {OUTPUT_XLSX.name} (colonne id forcée au format Texte)")
    print("\nPour Excel, utilisez de préférence id_output.xlsx.")
    return 0


if __name__ == "__main__":
    try:
        exit_code = main()
    except KeyboardInterrupt:
        print("\n\nOpération annulée.")
        exit_code = 1

    if sys.platform.startswith("win"):
        input("\nAppuyez sur Entrée pour fermer...")

    raise SystemExit(exit_code)
